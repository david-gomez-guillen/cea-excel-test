"""Run an Excel workbook from code, as a function from input cells to output cells.

The workbook is compiled once with the `formulas` package (a Python
implementation of Excel's calculation engine). Each run writes values into the
input cells, recalculates and reads the output cells, so the logic lives only in
the xlsx and nothing here duplicates it. The file on disk is never modified.

Which cells are inputs and outputs is declared in an interface table, one row
per input or output:

    Name          Role    Cell              Label   Unit   Group   Min  Max  Labels
    p_cure        input   Inputs!B13        ...
    p_onset       input                     ...                             age_groups
    cost          output  Results!C8:C10    ...

  Name    What code calls it.
  Role    input or output.
  Cell    A cell, a range or a defined name. Blank: the defined name called Name.
  Label, Unit, Group
          Text that describes it (optional).
  Min, Max
          Bounds of an input. Blank: those of the cell's data validation, if any.
  Labels  A range or defined name with one label per row of a table (per cell if
          it is a single row), such as the age groups of an age table (optional).

The table is read from a sheet named Interface in the workbook, so it travels
with the file and can be edited in Excel. Its labels and units can be formulas
pointing to the workbook's own text, which is read from the values saved with
the file. For a workbook that should not be edited, the same table can be a CSV
file or a list of dicts.

    from xlsx_model import XlsxModel

    model = XlsxModel('model.xlsx')       # or XlsxModel('model.xlsx', 'interface.csv')
    model.inputs['p_cure'].base           # 0.9, the value saved in the file
    results = model.run({'p_cure': 0.8})  # {'cost': [1203.5, 1544.0, 1790.2]}

Values are a number or text for a single cell, a list for a single row or
column, and a list of rows for a table. A table input can also be given as a dict
keyed by its labels. Inputs not given keep the values saved in the file.
"""

import csv
import logging
from dataclasses import dataclass, field
from pathlib import Path

import formulas
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

INTERFACE_SHEET = 'Interface'
COLUMNS = ('name', 'role', 'cell', 'label', 'unit', 'group', 'min', 'max', 'labels')


@dataclass
class Variable:
    """An input or output of the workbook: where it is and what it is."""
    name: str
    role: str
    sheet: str
    cells: list            # coordinates, row by row
    rows: int
    cols: int
    label: str = None
    unit: str = None
    group: str = None
    min: float = None
    max: float = None
    labels: list = None    # one per row, or per cell for a single row
    base: object = field(default=None, repr=False)  # the value saved in the file

    def shaped(self, values):
        """Flat cell values -> a value, a list or a list of rows."""
        if len(values) == 1:
            return values[0]
        if self.rows == 1 or self.cols == 1:
            return list(values)
        return [list(values[r * self.cols:(r + 1) * self.cols]) for r in range(self.rows)]


def _flat(value):
    """Nested lists, tuples and arrays -> a flat list of Python values."""
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, (list, tuple)):
        return [x for v in value for x in _flat(v)]
    return [value.item() if isinstance(value, np.generic) else value]


def _number(text):
    if text is None or (isinstance(text, str) and not text.strip()):
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None  # a bound given by a formula or a reference


def _cell_value(value, where):
    v = np.asarray(getattr(value, 'value', value)).ravel()[0]
    if isinstance(v, np.generic):
        v = v.item()
    if isinstance(v, str) and v.startswith('#'):
        raise ValueError(f'The workbook returned the error {v} in {where}')
    return None if v == '' else v


def read_interface(source):
    """The rows of an interface table, as dicts with the keys of COLUMNS.

    source is a worksheet, the path of a CSV file or a list of dicts. In a
    worksheet the table starts at the row with the headers Name and Role, so the
    sheet can explain itself above it. Headers are matched without case and only
    on their first word, so "Min (blank: from the data validation)" is Min.
    """
    if isinstance(source, (str, Path)):
        with open(source, newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
    elif hasattr(source, 'iter_rows'):
        rows = [list(r) for r in source.iter_rows(values_only=True)]
    else:
        rows = None
    if rows is not None:
        def key(h):
            return str(h).strip().split()[0].lower() if h is not None and str(h).strip() else None
        start = next((i for i, r in enumerate(rows) if {'name', 'role'} <= {key(h) for h in r}), None)
        if start is None:
            raise ValueError('The interface table needs a header row with at least Name and Role')
        headers = [key(h) for h in rows[start]]
        source = [dict(zip(headers, r)) for r in rows[start + 1:]]
    table = []
    for row in source:
        row = {str(k).lower(): v for k, v in row.items() if k is not None}
        row = {k: (v.strip() if isinstance(v, str) else v) for k, v in row.items()}
        row = {k: (None if v == '' else v) for k, v in row.items()}
        if row.get('name') is None:
            continue
        unknown = set(row) - set(COLUMNS)
        if unknown:
            raise ValueError(f'Unknown interface columns: {", ".join(sorted(unknown))}')
        table.append({c: row.get(c) for c in COLUMNS})
    return table


class XlsxModel:

    def __init__(self, path, interface=None):
        """path: the workbook. interface: None to read the Interface sheet of the
        workbook, or the path of a CSV file, or a list of dicts."""
        self.path = Path(path)
        # The values saved with the file: the inputs' base values and the text of
        # the interface table, even where it is a formula.
        self._wb = openpyxl.load_workbook(self.path, data_only=True)
        if interface is None:
            if INTERFACE_SHEET not in self._wb.sheetnames:
                raise ValueError(f'{self.path.name} has no {INTERFACE_SHEET} sheet: '
                                 'pass the interface table as a CSV file or a list of dicts')
            interface = self._wb[INTERFACE_SHEET]
        self._bounds = {}
        self.inputs, self.outputs = {}, {}
        for row in read_interface(interface):
            var = self._variable(row)
            if var.name in self.inputs or var.name in self.outputs:
                raise ValueError(f'{var.name} is declared twice in the interface')
            (self.inputs if var.role == 'input' else self.outputs)[var.name] = var
        if not self.outputs:
            raise ValueError('The interface declares no outputs')
        del self._wb

        logging.getLogger('formulas').setLevel(logging.WARNING)
        # Compile the workbook into a function of every input cell that returns
        # every output cell: much faster than recalculating the whole workbook.
        self._base_cells = {self._key(v.sheet, c): x for v in self.inputs.values()
                            for c, x in zip(v.cells, _flat(v.base))}
        self._input_keys = list(self._base_cells)
        self._output_keys = [self._key(v.sheet, c) for v in self.outputs.values() for c in v.cells]
        model = formulas.ExcelModel().loads(str(self.path)).finish()
        self._function = model.compile(inputs=self._input_keys, outputs=self._output_keys)

    def _key(self, sheet, cell):
        return f"'[{self.path.name}]{sheet.upper()}'!{cell}"

    def _range(self, ref, what):
        """'Inputs'!$C$43:$C$54 or a defined name -> (sheet, cells row by row, rows, cols)."""
        ref = str(ref).lstrip('=')
        if '!' not in ref:
            if ref not in self._wb.defined_names:
                raise ValueError(f'{what}: no cell given and no defined name called {ref}')
            ref = self._wb.defined_names[ref].attr_text
        if ',' in ref:
            raise ValueError(f'{what}: {ref} has several areas; use a single rectangular range')
        sheet, rng = ref.rsplit('!', 1)
        sheet = sheet.strip("'").replace("''", "'")
        if sheet not in self._wb.sheetnames:
            raise ValueError(f'{what}: there is no sheet called {sheet}')
        min_col, min_row, max_col, max_row = range_boundaries(rng.replace('$', ''))
        cells = [f'{get_column_letter(c)}{r}' for r in range(min_row, max_row + 1)
                 for c in range(min_col, max_col + 1)]
        return sheet, cells, max_row - min_row + 1, max_col - min_col + 1

    def _validation_bounds(self, sheet):
        """Cell -> (min, max) of the numeric data validations of a sheet."""
        if sheet not in self._bounds:
            bounds = {}
            for dv in self._wb[sheet].data_validations.dataValidation:
                if dv.type not in ('decimal', 'whole'):
                    continue
                a, b = _number(dv.formula1), _number(dv.formula2)
                lo, hi = {'between': (a, b), 'equal': (a, a),
                          'greaterThan': (a, None), 'greaterThanOrEqual': (a, None),
                          'lessThan': (None, a), 'lessThanOrEqual': (None, a)
                          }.get(dv.operator or 'between', (None, None))
                for rng in str(dv.sqref).split():
                    for row in self._wb[sheet][rng] if ':' in rng else [[self._wb[sheet][rng]]]:
                        for cell in row:
                            bounds[cell.coordinate] = (lo, hi)
            self._bounds[sheet] = bounds
        return self._bounds[sheet]

    def _variable(self, row):
        name, role = str(row['name']), str(row['role'] or '').lower()
        if role not in ('input', 'output'):
            raise ValueError(f'{name}: Role must be input or output, not {row["role"]!r}')
        sheet, cells, rows, cols = self._range(row['cell'] or name, name)
        var = Variable(name, role, sheet, cells, rows, cols, label=row['label'], unit=row['unit'],
                       group=row['group'])
        ws = self._wb[sheet]
        var.base = var.shaped([ws[c].value for c in cells])
        if row['labels'] is not None:
            lsheet, lcells, *_ = self._range(row['labels'], f'{name} (labels)')
            var.labels = [self._wb[lsheet][c].value for c in lcells]
            expected = cols if rows == 1 else rows
            if len(var.labels) != expected:
                raise ValueError(f'{name} needs {expected} labels, one per '
                                 f'{"cell" if rows == 1 else "row"}, got {len(var.labels)}')
        if role == 'input':
            # Bounds every cell meets: the loosest of their data validations.
            bounds = [self._validation_bounds(sheet).get(c, (None, None)) for c in cells]
            los, his = [b[0] for b in bounds], [b[1] for b in bounds]
            var.min = _number(row['min']) if row['min'] is not None else (
                None if None in los else min(los))
            var.max = _number(row['max']) if row['max'] is not None else (
                None if None in his else max(his))
        return var

    def _input_cells(self, var, value):
        if isinstance(value, dict):
            if var.labels is None:
                raise ValueError(f'{var.name} has no labels to key its values by')
            missing = [lab for lab in var.labels if lab not in value]
            extra = [k for k in value if k not in var.labels]
            if missing or extra:
                raise ValueError(f'{var.name} needs one value per label: '
                                 f'missing {missing}, unknown {extra}')
            value = [value[lab] for lab in var.labels]
        values = _flat(value)
        if len(values) != len(var.cells):
            raise ValueError(f'{var.name} needs {len(var.cells)} values, got {len(values)}')
        return zip([self._key(var.sheet, c) for c in var.cells], values)

    def run(self, values=None):
        """Recalculates the workbook with the given inputs, {name: value}, and
        returns every output, {name: value}. Blank cells are None."""
        cells = dict(self._base_cells)
        for name, value in (values or {}).items():
            if name not in self.inputs:
                raise KeyError(f'Unknown input: {name}')
            cells.update(self._input_cells(self.inputs[name], value))
        sol = self._function(*[cells[k] for k in self._input_keys])
        # With a single output cell the compiled function returns it alone.
        sol = dict(zip(self._output_keys, [sol] if len(self._output_keys) == 1 else sol))
        return {name: var.shaped([_cell_value(sol[self._key(var.sheet, c)], f'{name} ({var.sheet}!{c})')
                                  for c in var.cells])
                for name, var in self.outputs.items()}

"""Run cea_model.xlsx from code, using the workbook itself as the model.

The workbook is compiled once with the `formulas` package (a Python
implementation of Excel's calculation engine). Each run writes the inputs into
the named input cells, recalculates and reads the named output cells, so the
logic lives only in the xlsx and nothing here duplicates it. The file on disk is
never modified.

    from cea_workbook import WorkbookModel

    model = WorkbookModel()
    results = model.run({'screen_participation': 0.8, 'wtp': 30000})
    for row in results['summary']:
        print(row['strategy'], row['cost'], row['qalys'])

Parameters are the input names of the workbook (column "Name in formulas" of the
Inputs sheet): a number for a single input, or a list of 12 numbers for the
age tables `p_background_death` and `p_cancer_onset`. Strategies can be
replaced with a list of up to three dicts with keys name, screens ('Yes' or
'No'), first_age and last_age.
"""

import logging
import re
from pathlib import Path

import formulas
import numpy as np
import openpyxl

DEFAULT_PATH = Path(__file__).with_name('cea_model.xlsx')

TABLE_INPUTS = ('p_background_death', 'p_cancer_onset')
STRATEGY_FIELDS = {'name': 'strategy_name', 'screens': 'strategy_screens',
                   'first_age': 'strategy_first_age', 'last_age': 'strategy_last_age'}
# Every named cell that is not an input or an output (lookup keys of the tables).
STRUCTURAL = ('age_band_label', 'age_band_start')


def _cells(ref):
    """'Inputs'!$C$37:$C$48 -> ('Inputs', ['C37', ..., 'C48'])."""
    sheet, rng = ref.rsplit('!', 1)
    sheet = sheet.strip("'")
    rng = rng.replace('$', '')
    if ':' not in rng:
        return sheet, [rng]
    min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(rng)
    return sheet, [f'{openpyxl.utils.get_column_letter(c)}{r}'
                   for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]


def _scalar(value):
    v = np.asarray(getattr(value, 'value', value)).ravel()[0]
    if isinstance(v, np.generic):
        v = v.item()
    if isinstance(v, str) and v.startswith('#'):
        raise ValueError(f'The workbook returned the error {v}')
    return v


class WorkbookModel:

    def __init__(self, path=DEFAULT_PATH):
        self.path = Path(path)
        wb = openpyxl.load_workbook(self.path)
        self.names = {nm: _cells(d.attr_text) for nm, d in wb.defined_names.items()}
        # The inputs are the named cells of the Inputs sheet that hold values
        # (not formulas, like its input check).
        self.input_names = [nm for nm, (sheet, cells) in self.names.items()
                            if sheet == 'Inputs' and nm not in STRUCTURAL
                            and not str(wb[sheet][cells[0]].value).startswith('=')]
        self.base_parameters = {}
        for nm in self.input_names:
            sheet, cells = self.names[nm]
            values = [wb[sheet][c].value for c in cells]
            self.base_parameters[nm] = values if len(values) > 1 else values[0]
        self.age_groups = [wb['Inputs'][c].value for c in self.names['age_band_label'][1]]
        self.parameter_info = self._read_parameter_info(wb)

        logging.getLogger('formulas').setLevel(logging.WARNING)
        self._book = self.path.name
        # Compile the workbook into a function of every input cell that returns
        # every output cell: much faster than recalculating the whole workbook.
        self._input_keys = list(self._inputs(self._base_values(), self.base_strategies))
        self._output_keys = [self._key(sheet, c) for nm, (sheet, cells) in self.names.items()
                             if nm.startswith('out_') for c in cells]
        model = formulas.ExcelModel().loads(str(self.path)).finish()
        self._function = model.compile(inputs=self._input_keys, outputs=self._output_keys)

    def _read_parameter_info(self, wb):
        """What the Inputs sheet says about each parameter.

        For each parameter: its label, unit, section (the heading it is under,
        without its number) and the bounds of its data validation, None when it
        has none.
        """
        ws = wb['Inputs']
        bounds = {}
        for dv in ws.data_validations.dataValidation:
            if dv.type not in ('decimal', 'whole'):
                continue
            lo = float(dv.formula1) if dv.formula1 is not None else None
            hi = float(dv.formula2) if dv.operator == 'between' and dv.formula2 is not None else None
            for rng in str(dv.sqref).split():
                for row in ws[rng] if ':' in rng else [[ws[rng]]]:
                    for cell in row:
                        bounds[cell.coordinate] = (lo, hi)
        info = {}
        for nm in self.parameter_names:
            _, cells = self.names[nm]
            first = ws[cells[0]]
            if nm in TABLE_INPUTS:
                label, unit = ws.cell(first.row - 1, first.column).value, 'per year'
            else:
                label, unit = ws.cell(first.row, 1).value, ws.cell(first.row, 3).value
            section = None
            for r in range(first.row - 1, 0, -1):
                m = re.match(r'^\d+\.\s*(.+)$', str(ws.cell(r, 1).value or ''))
                if m:
                    section = m.group(1)
                    break
            lo, hi = bounds.get(cells[0], (None, None))
            info[nm] = {'label': label, 'unit': unit, 'section': section, 'min': lo, 'max': hi}
        return info

    def _base_values(self):
        return {nm: self.base_parameters[nm] for nm in self.parameter_names}

    @property
    def parameter_names(self):
        return [nm for nm in self.input_names if not nm.startswith('strategy_')]

    @property
    def base_strategies(self):
        b = self.base_parameters
        return [{field: b[nm][i] for field, nm in STRATEGY_FIELDS.items()}
                for i in range(len(b['strategy_name']))]

    def _key(self, sheet, cell):
        return f"'[{self._book}]{sheet.upper()}'!{cell}"

    def _inputs(self, parameters, strategies):
        values = {}
        for nm, value in (parameters or {}).items():
            if nm not in self.input_names or nm.startswith('strategy_'):
                raise KeyError(f'Unknown parameter: {nm}')
            sheet, cells = self.names[nm]
            value = list(np.ravel(value)) if nm in TABLE_INPUTS else [value]
            if len(value) != len(cells):
                raise ValueError(f'{nm} needs {len(cells)} values, got {len(value)}')
            for cell, v in zip(cells, value):
                values[self._key(sheet, cell)] = v
        if strategies is not None:
            n = len(self.base_parameters['strategy_name'])
            if not 1 <= len(strategies) <= n:
                raise ValueError(f'The workbook compares 1 to {n} strategies')
            # Unused slots repeat the first strategy; their results are dropped.
            padded = list(strategies) + [strategies[0]] * (n - len(strategies))
            for field, nm in STRATEGY_FIELDS.items():
                sheet, cells = self.names[nm]
                for cell, s in zip(cells, padded):
                    values[self._key(sheet, cell)] = s[field]
        return values

    def run(self, parameters=None, strategies=None):
        """Recalculates the workbook with the given inputs and returns its results.

        Returns a dict with:
          summary:   one dict per strategy: cost, qalys, life_years and
                     life_years_undiscounted (per person, discounted unless said
                     otherwise), nmb, icer_vs_first, cancer_deaths_per_100k,
                     diagnosed_early_per_100k, diagnosed_advanced_per_100k.
          frontier:  the strategies sorted by cost, with status and ICER.
          incidence: per strategy, cancers diagnosed per 100,000 person-years
                     for each age group (None for groups the cohort never reaches).
          age_groups, best (the conclusion sentence) and check ('All OK' if the
          internal checks of the workbook pass).
        """
        values = self._inputs(self._base_values(), self.base_strategies)
        values.update(self._inputs(parameters, strategies))
        sol = dict(zip(self._output_keys, self._function(*[values[k] for k in self._input_keys])))

        def get(nm):
            sheet, cells = self.names[nm]
            out = []
            for c in cells:
                v = _scalar(sol[self._key(sheet, c)])
                out.append(None if v == '' else v)
            return out

        n = len(strategies) if strategies is not None else len(self.base_parameters['strategy_name'])
        names = get('out_strategy')[:n]
        deaths = get('out_cancer_deaths')
        diagnoses = get('out_diagnoses')  # early, advanced per strategy
        columns = {'cost': get('out_cost'), 'qalys': get('out_qalys'),
                   'life_years': get('out_life_years'),
                   'life_years_undiscounted': get('out_life_years_undiscounted'),
                   'nmb': get('out_nmb'), 'icer_vs_first': get('out_icer_vs_first')}
        summary = []
        for i, s in enumerate(names):
            row = {'strategy': s}
            row.update({k: v[i] for k, v in columns.items()})
            row['cancer_deaths_per_100k'] = deaths[i]
            row['diagnosed_early_per_100k'] = diagnoses[2 * i]
            row['diagnosed_advanced_per_100k'] = diagnoses[2 * i + 1]
            summary.append(row)

        incidence_cells = get('out_incidence')
        n_slots = len(self.base_parameters['strategy_name'])
        incidence = {s: incidence_cells[i::n_slots] for i, s in enumerate(names)}

        frontier = []
        if n == n_slots:
            frontier = [{'strategy': s, 'status': st, 'icer': None if ic == '—' else ic}
                        for s, st, ic in zip(get('out_frontier_strategy'), get('out_frontier_status'),
                                             [None] + get('out_frontier_icer')[1:])]
        return {'summary': summary, 'frontier': frontier, 'incidence': incidence,
                'age_groups': self.age_groups, 'best': get('out_best')[0],
                'check': get('out_check')[0]}


if __name__ == '__main__':
    import time
    t = time.time()
    model = WorkbookModel()
    print(f'Compiled in {time.time() - t:.1f} s')
    t = time.time()
    res = model.run()
    print(f'Ran in {time.time() - t:.2f} s')
    print(res['check'], '|', res['best'])
    for row in res['summary']:
        print(row)
    for row in res['frontier']:
        print(row)
